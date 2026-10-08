#!/usr/bin/env python3
"""
Juniper Mist — Organization Switch Port Inventory

Prompts for a Mist API token (hidden via getpass), selects or auto-detects the
cloud region, runs a self-check against GET /api/v1/self, then exports every
site, every switch (including Virtual Chassis members), and every port.

Pagination follows Mist's own rules:
  * list endpoints continue until every row counted by X-Page-Total is collected
  * search endpoints follow the `next` URL (the search_after cursor lives in
    that URL and must not be rebuilt by hand)

Typical usage:
    python3 mist_switch_port_inventory.py

Optional:
    python3 mist_switch_port_inventory.py --host api.mist.com
    MIST_API_TOKEN=... python3 mist_switch_port_inventory.py --no-prompt-token

Dependencies:
    pip install requests openpyxl
"""

from __future__ import annotations

import argparse
import getpass
import os
import re
import sys
import time
from datetime import datetime, timezone
from typing import Any

try:
    import requests
except ImportError:
    sys.exit("Missing dependency: requests\n  pip install requests openpyxl")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    sys.exit("Missing dependency: openpyxl\n  pip install requests openpyxl")


MIST_REGIONS: list[dict[str, str]] = [
    {"name": "Global 01", "host": "api.mist.com"},
    {"name": "Global 02", "host": "api.gc1.mist.com"},
    {"name": "Global 03", "host": "api.ac2.mist.com"},
    {"name": "Global 04", "host": "api.gc2.mist.com"},
    {"name": "Global 05", "host": "api.gc4.mist.com"},
    {"name": "EMEA 01", "host": "api.eu.mist.com"},
    {"name": "EMEA 02", "host": "api.gc3.mist.com"},
    {"name": "EMEA 03", "host": "api.ac6.mist.com"},
    {"name": "EMEA 04", "host": "api.gc6.mist.com"},
    {"name": "APAC 01", "host": "api.ac5.mist.com"},
    {"name": "APAC 02", "host": "api.gc5.mist.com"},
    {"name": "APAC 03", "host": "api.gc7.mist.com"},
]

DEFAULT_TIMEOUT = 60
PAGE_LIMIT = 1000
RATE_SLEEP = 0.05
MAX_PAGES = 2000


class MistAPIError(RuntimeError):
    """Raised when a Mist API call fails in a way the script cannot continue."""


def norm_mac(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[^0-9a-f]", "", str(value).lower())


def norm_port(value: Any) -> str:
    return str(value or "").strip().lower()


def _header_int(headers: Any, name: str) -> int | None:
    if not headers:
        return None
    raw = headers.get(name)
    if raw is None or raw == "":
        return None
    try:
        return int(float(str(raw).strip()))
    except (TypeError, ValueError):
        return None


def _encode_params(params: dict[str, Any] | None) -> dict[str, Any]:
    """Mist bools are the lowercase strings true/false. requests would send True."""
    encoded: dict[str, Any] = {}
    for key, value in (params or {}).items():
        if value is None:
            continue
        if isinstance(value, bool):
            encoded[key] = "true" if value else "false"
        else:
            encoded[key] = value
    return encoded


def _fingerprint(batch: list[Any]) -> str:
    def ident(item: Any) -> str:
        if isinstance(item, dict):
            return str(
                item.get("mac")
                or item.get("id")
                or item.get("port_id")
                or item.get("serial")
                or item.get("name")
                or ""
            )
        return str(item)

    if not batch:
        return ""
    mid = batch[len(batch) // 2]
    return f"{len(batch)}|{ident(batch[0])}|{ident(mid)}|{ident(batch[-1])}"


def _retry_wait(resp: requests.Response, attempt: int) -> float:
    raw = resp.headers.get("Retry-After") if resp is not None else None
    if raw:
        try:
            return min(float(raw), 120.0)
        except ValueError:
            pass
    return min(2 ** attempt, 30)


class MistClient:
    def __init__(self, host: str, token: str, timeout: int = DEFAULT_TIMEOUT) -> None:
        self.host = host.removeprefix("https://").removeprefix("http://").rstrip("/")
        self.base = f"https://{self.host}/api/v1"
        self.timeout = timeout
        self.session = requests.Session()
        self.session.headers.update(
            {
                "Authorization": f"Token {token}",
                "Accept": "application/json",
                "Content-Type": "application/json",
                "User-Agent": "mist-switch-port-inventory/1.1",
            }
        )

    def resolve_url(self, path: str) -> str:
        path = str(path).strip()
        if path.startswith(("http://", "https://")):
            return path
        if path.startswith("/api/"):
            return f"https://{self.host}{path}"
        if not path.startswith("/"):
            path = "/" + path
        return f"{self.base}{path}"

    def request(self, path: str, params: dict[str, Any] | None = None) -> tuple[Any, Any]:
        url = self.resolve_url(path)
        query = _encode_params(params)
        extra_headers: dict[str, str] = {}
        if query and "page" in query:
            extra_headers["X-Page-Limit"] = str(query.get("limit", PAGE_LIMIT))
            extra_headers["X-Page-Page"] = str(query["page"])

        last_error: Exception | None = None
        for attempt in range(6):
            try:
                resp = self.session.get(
                    url, params=query or None, headers=extra_headers, timeout=self.timeout
                )
            except requests.RequestException as exc:
                last_error = exc
                wait = min(2 ** attempt, 30)
                print(f"  Network error calling {url}: {exc} — retry in {wait:.0f}s")
                time.sleep(wait)
                continue

            if resp.status_code == 429 or resp.status_code >= 500:
                wait = _retry_wait(resp, attempt)
                print(f"  HTTP {resp.status_code} from {path} — retry in {wait:.0f}s")
                time.sleep(wait)
                last_error = MistAPIError(f"HTTP {resp.status_code} from {path}")
                continue

            if resp.status_code == 401:
                raise MistAPIError(
                    f"401 Unauthorized from {self.host}{path}. "
                    "Token is invalid for this cloud, expired, or missing privileges."
                )
            if resp.status_code == 403:
                raise MistAPIError(
                    f"403 Forbidden from {self.host}{path}. "
                    "Token does not have permission for this resource."
                )
            if resp.status_code >= 400:
                detail = resp.text[:400]
                raise MistAPIError(f"HTTP {resp.status_code} from {path}: {detail}")

            if not resp.content:
                return None, resp.headers
            try:
                return resp.json(), resp.headers
            except ValueError as exc:
                raise MistAPIError(f"Non-JSON response from {path}: {resp.text[:200]}") from exc

        raise MistAPIError(f"Giving up on {url}: {last_error}")

    def get(self, path: str, params: dict[str, Any] | None = None) -> Any:
        data, _headers = self.request(path, params)
        return data

    def get_paginated(self, path: str, params: dict[str, Any] | None = None) -> list[Any]:
        """Walk list endpoints until X-Page-Total rows have been collected."""
        params = dict(params or {})
        params.setdefault("limit", PAGE_LIMIT)
        requested = int(params["limit"])
        page = 1
        items: list[Any] = []
        previous_fp = None
        while page <= MAX_PAGES:
            params["page"] = page
            data, headers = self.request(path, params)
            if isinstance(data, dict) and ("results" in data or data.get("next")):
                search_params = {k: v for k, v in params.items() if k != "page"}
                return self.search_paginated(path, search_params)
            if not isinstance(data, list):
                if data and not items:
                    return [data]
                return items
            if not data:
                break
            fp = _fingerprint(data)
            if fp == previous_fp:
                print(f"  Stopped {path}: page {page} repeated the previous page")
                break
            previous_fp = fp
            items.extend(data)

            total = _header_int(headers, "X-Page-Total")
            # Do not stop just because this page is shorter than the limit we
            # asked for. Mist often returns 100 rows while echoing limit=1000,
            # with the real count only in X-Page-Total.
            if total is not None:
                if len(items) >= total:
                    break
            elif len(data) < requested:
                break
            page += 1
            if page == 2 or page % 5 == 0:
                print(f"    {path} page {page - 1}: {len(items)}/{total or '?'} so far")
            time.sleep(RATE_SLEEP)
        else:
            print(f"  Warning: stopped {path} after {MAX_PAGES} pages")
        return items

    def search_paginated(self, path: str, params: dict[str, Any] | None = None) -> list[Any]:
        """Walk search endpoints by GETting each response's `next` URL."""
        params = {k: v for k, v in dict(params or {}).items() if k != "page"}
        params.setdefault("limit", PAGE_LIMIT)
        items: list[Any] = []
        seen_next: set[str] = set()
        next_url: str | None = None
        first = True
        pages = 0
        while pages < MAX_PAGES:
            pages += 1
            if first:
                data, _headers = self.request(path, params)
                first = False
            else:
                data, _headers = self.request(next_url or path)
            if isinstance(data, list):
                items.extend(data)
                break
            if not isinstance(data, dict):
                break
            batch = data.get("results")
            if batch is None and isinstance(data.get("data"), list):
                batch = data.get("data")
            if not isinstance(batch, list):
                if not items and data and "results" not in data:
                    return [data]
                break
            if not batch:
                break
            items.extend(batch)
            try:
                total_i = int(data["total"]) if data.get("total") is not None else None
            except (TypeError, ValueError):
                total_i = None
            if total_i is not None and len(items) >= total_i:
                break
            nxt = data.get("next")
            if not nxt:
                if total_i is not None and len(items) < total_i:
                    print(
                        f"  Warning: {path} returned {len(items)} of {total_i} "
                        "and no next cursor"
                    )
                break
            nxt = str(nxt).strip()
            if not nxt or nxt in seen_next:
                break
            seen_next.add(nxt)
            next_url = nxt
            if pages == 1 or pages % 5 == 0:
                print(f"    {path} search page {pages}: {len(items)}/{total_i or '?'}")
            time.sleep(RATE_SLEEP)
        else:
            print(f"  Warning: stopped {path} after {MAX_PAGES} search pages")
        return items


def prompt_token(cli_token: str | None, no_prompt: bool) -> str:
    token = (cli_token or os.environ.get("MIST_API_TOKEN") or "").strip()
    if token:
        return token
    if no_prompt:
        sys.exit("No API token provided. Set MIST_API_TOKEN or omit --no-prompt-token.")
    print()
    print("Create a user token at  https://{api-host}/api/v1/self/apitokens")
    print("or an org token under Organization > Admin > Settings > API Token.")
    print("Send it as:  Authorization: Token <value>")
    print()
    token = getpass.getpass("Mist API token (input hidden): ").strip()
    if not token:
        sys.exit("No API token entered.")
    return token


def choose_region(client_factory, token: str, forced_host: str | None) -> tuple[MistClient, dict[str, str]]:
    if forced_host:
        host = forced_host.removeprefix("https://").removeprefix("http://").rstrip("/")
        region = next((r for r in MIST_REGIONS if r["host"] == host), {"name": host, "host": host})
        client = client_factory(host, token)
        client.get("/self")
        return client, region

    print()
    print("Mist API tokens are bound to one cloud region and the region is not")
    print("encoded in the token. Choose a region or let the script probe all clouds.")
    print()
    print("  0) Auto-detect (try every region until /api/v1/self succeeds)")
    for idx, region in enumerate(MIST_REGIONS, start=1):
        print(f"  {idx}) {region['name']:<12}  {region['host']}")
    print()

    while True:
        raw = input(f"Select region [0-{len(MIST_REGIONS)}] (default 0): ").strip() or "0"
        if raw.isdigit() and 0 <= int(raw) <= len(MIST_REGIONS):
            choice = int(raw)
            break
        print("  Invalid selection.")

    if choice > 0:
        region = MIST_REGIONS[choice - 1]
        print(f"\nUsing {region['name']} ({region['host']})")
        client = client_factory(region["host"], token)
        client.get("/self")
        return client, region

    print("\nProbing regional API endpoints with GET /api/v1/self ...")
    last_error = None
    for region in MIST_REGIONS:
        try:
            client = client_factory(region["host"], token)
            client.get("/self")
            print(f"  [ok]   {region['name']:<12} {region['host']}")
            print(f"\nDetected region: {region['name']} ({region['host']})")
            return client, region
        except MistAPIError as exc:
            last_error = exc
            msg = str(exc)
            if "401" in msg or "403" in msg:
                print(f"  [auth] {region['name']:<12} {region['host']} — token rejected")
            else:
                print(f"  [miss] {region['name']:<12} {region['host']}")
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            print(f"  [miss] {region['name']:<12} {region['host']} — {exc}")

    sys.exit(
        "Could not authenticate against any known Mist cloud.\n"
        f"Last error: {last_error}\n"
        "Check the token and that it was created in the same cloud as the org."
    )


def self_check(client: MistClient) -> dict[str, Any]:
    print("\n=== Self-check  GET /api/v1/self ===")
    me = client.get("/self") or {}
    if not isinstance(me, dict):
        raise MistAPIError("Unexpected /self payload")

    email = me.get("email") or me.get("name") or "(org token / unnamed)"
    first = me.get("first_name") or ""
    last = me.get("last_name") or ""
    display = " ".join(p for p in (first, last) if p) or email
    print(f"  Account     : {display}")
    print(f"  Email       : {email}")
    if me.get("via_sso") is not None:
        print(f"  SSO         : {me.get('via_sso')}")
    tags = me.get("tags")
    if tags:
        print(f"  Tags        : {tags}")

    privileges = me.get("privileges") or []
    print(f"  Privileges  : {len(privileges)}")
    for priv in privileges:
        scope = priv.get("scope", "?")
        role = priv.get("role", "?")
        name = priv.get("name") or priv.get("org_name") or ""
        org_id = priv.get("org_id") or ""
        site = priv.get("site_id") or priv.get("sitegroup_id") or ""
        extra = f" site={site}" if site else ""
        print(f"    - {scope:<8} role={role:<12} {name}  org={org_id}{extra}")

    if not privileges and not me.get("email"):
        print("  Note: org tokens often return a slim /self body. Continuing.")
    return me


def orgs_from_self(me: dict[str, Any]) -> list[dict[str, str]]:
    seen: dict[str, dict[str, str]] = {}
    for priv in me.get("privileges") or []:
        org_id = priv.get("org_id")
        if not org_id:
            continue
        name = priv.get("name") if priv.get("scope") == "org" else priv.get("org_name")
        entry = seen.setdefault(org_id, {"org_id": org_id, "name": name or org_id, "roles": []})
        if name and entry["name"] == org_id:
            entry["name"] = name
        role = priv.get("role")
        if role and role not in entry["roles"]:
            entry["roles"].append(role)
    return list(seen.values())


def choose_org(orgs: list[dict[str, str]], client: MistClient) -> dict[str, str]:
    if not orgs:
        raw = input("\nNo org found on /self. Enter organization UUID: ").strip()
        if not raw:
            sys.exit("Organization ID is required.")
        info = {"org_id": raw, "name": raw, "roles": []}
        try:
            org = client.get(f"/orgs/{raw}") or {}
            info["name"] = org.get("name") or raw
        except MistAPIError as exc:
            print(f"  Warning: could not fetch org details ({exc})")
        return info

    if len(orgs) == 1:
        org = orgs[0]
        print(f"\nUsing organization: {org['name']}  ({org['org_id']})")
        return org

    print("\nOrganizations this token can access:")
    for idx, org in enumerate(orgs, start=1):
        roles = ",".join(org.get("roles") or []) or "-"
        print(f"  {idx}) {org['name']}  [{roles}]  {org['org_id']}")
    while True:
        raw = input(f"Select organization [1-{len(orgs)}]: ").strip()
        if raw.isdigit() and 1 <= int(raw) <= len(orgs):
            return orgs[int(raw) - 1]
        print("  Invalid selection.")


def switch_key(sw: dict[str, Any]) -> str:
    return norm_mac(sw.get("mac")) or str(sw.get("id") or sw.get("serial") or "")


def merge_switch(store: dict[str, dict[str, Any]], incoming: dict[str, Any]) -> None:
    key = switch_key(incoming)
    if not key:
        return
    current = store.get(key)
    if current is None:
        copied = dict(incoming)
        if copied.get("mac"):
            copied["mac"] = norm_mac(copied.get("mac")) or copied.get("mac")
        if copied.get("vc_mac"):
            copied["vc_mac"] = norm_mac(copied.get("vc_mac")) or copied.get("vc_mac")
        store[key] = copied
        return
    for field, value in incoming.items():
        if value in (None, "", [], {}):
            continue
        if field in ("mac", "vc_mac"):
            value = norm_mac(value) or value
        if current.get(field) in (None, "", [], {}):
            current[field] = value


def _as_switch_record(raw: dict[str, Any], site_id: str | None = None) -> dict[str, Any]:
    status = raw.get("status")
    connected = raw.get("connected")
    if connected is None and isinstance(status, str):
        connected = status.lower() == "connected"
    return {
        "id": raw.get("id") or raw.get("device_id") or "",
        "mac": raw.get("mac") or "",
        "vc_mac": raw.get("vc_mac") or raw.get("chassis_mac") or "",
        "name": raw.get("name") or raw.get("hostname") or "",
        "hostname": raw.get("hostname") or "",
        "model": raw.get("model") or "",
        "serial": raw.get("serial") or "",
        "sku": raw.get("sku") or "",
        "site_id": raw.get("site_id") or site_id or "",
        "connected": connected,
        "status": status if isinstance(status, str) else "",
        "adopted": raw.get("adopted"),
        "version": raw.get("version") or "",
        "vc_role": raw.get("vc_role") or raw.get("role") or "",
        "member_id": raw.get("member_id", raw.get("fpc", raw.get("idx"))),
        "type": raw.get("type") or "switch",
    }


def _member_record(parent: dict[str, Any], member: dict[str, Any], index: int) -> dict[str, Any]:
    parent_mac = norm_mac(parent.get("mac"))
    vc_mac = norm_mac(parent.get("vc_mac")) or parent_mac
    fpc = member.get("member_id", member.get("fpc", member.get("idx", index)))
    name = member.get("name") or ""
    if not name:
        base = parent.get("name") or parent.get("hostname") or vc_mac or "vc"
        name = f"{base} (fpc {fpc})"
    return {
        "id": member.get("id") or "",
        "mac": member.get("mac") or "",
        "vc_mac": vc_mac,
        "name": name,
        "hostname": member.get("hostname") or "",
        "model": member.get("model") or "",
        "serial": member.get("serial") or "",
        "sku": member.get("sku") or "",
        "site_id": member.get("site_id") or parent.get("site_id") or "",
        "connected": parent.get("connected"),
        "status": parent.get("status") or "",
        "adopted": parent.get("adopted"),
        "version": parent.get("version") or "",
        "vc_role": member.get("vc_role") or member.get("role") or "",
        "member_id": fpc,
        "type": "switch",
    }


def is_secondary_vc_member(sw: dict[str, Any]) -> bool:
    mac = norm_mac(sw.get("mac"))
    vc = norm_mac(sw.get("vc_mac"))
    return bool(mac and vc and mac != vc)


def fetch_sites(client: MistClient, org_id: str) -> list[dict[str, Any]]:
    print("\n=== Sites  GET /api/v1/orgs/{org_id}/sites ===")
    sites = client.get_paginated(f"/orgs/{org_id}/sites")
    seen: set[str] = set()
    unique: list[dict[str, Any]] = []
    for site in sites:
        sid = str(site.get("id") or "")
        if sid and sid in seen:
            continue
        if sid:
            seen.add(sid)
        unique.append(site)
    print(f"  Found {len(unique)} site(s)")
    for site in unique:
        print(f"    - {site.get('name', '(unnamed)')}  [{site.get('id')}]")
    return unique


def ensure_site(client: MistClient, sites: list[dict[str, Any]], site_id: str) -> None:
    if not site_id or any(s.get("id") == site_id for s in sites):
        return
    try:
        site = client.get(f"/sites/{site_id}") or {}
    except MistAPIError as exc:
        print(f"  Warning: could not load site {site_id} ({exc})")
        sites.append({"id": site_id, "name": site_id})
        return
    if isinstance(site, dict):
        site.setdefault("id", site_id)
        sites.append(site)


def harvest_port_config(
    cfg: dict[str, Any],
    mac: str,
    mapping: dict[tuple[str, str], dict[str, Any]],
) -> bool:
    port_config = cfg.get("port_config") or {}
    if not isinstance(port_config, dict) or not port_config:
        return False
    port_usages = cfg.get("port_usages") or {}
    if not isinstance(port_usages, dict):
        port_usages = {}
    for port_key, pcfg in port_config.items():
        if not isinstance(pcfg, dict):
            continue
        usage = pcfg.get("usage") or pcfg.get("dynamic_usage") or ""
        usage_def = port_usages.get(usage) if isinstance(port_usages.get(usage), dict) else {}
        networks = usage_def.get("networks")
        if isinstance(networks, list):
            networks_txt = ",".join(str(n) for n in networks)
        else:
            networks_txt = networks or ""
        record = {
            "usage": usage,
            "description": pcfg.get("description") or usage_def.get("description") or "",
            "disabled": pcfg.get("disabled"),
            "speed": pcfg.get("speed") or usage_def.get("speed"),
            "duplex": pcfg.get("duplex") or usage_def.get("duplex"),
            "poe_disabled": pcfg.get("poe_disabled")
            if "poe_disabled" in pcfg
            else usage_def.get("poe_disabled"),
            "mode": usage_def.get("mode") or pcfg.get("mode"),
            "port_network": usage_def.get("port_network") or pcfg.get("port_network"),
            "networks": networks_txt,
            "port_auth": usage_def.get("port_auth") or pcfg.get("port_auth"),
        }
        for expanded in expand_port_key(str(port_key)):
            mapping[(mac, norm_port(expanded))] = record
    return True


def fetch_switches(
    client: MistClient,
    org_id: str,
    sites: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[tuple[str, str], dict[str, Any]]]:
    print("\n=== Switches ===")
    store: dict[str, dict[str, Any]] = {}
    configs: dict[tuple[str, str], dict[str, Any]] = {}

    print("  GET /api/v1/orgs/{org_id}/inventory?type=switch&vc=true&unassigned=true")
    try:
        listed = client.get_paginated(
            f"/orgs/{org_id}/inventory",
            params={"type": "switch", "vc": True, "unassigned": True, "limit": PAGE_LIMIT},
        )
    except MistAPIError as exc:
        print(f"  Inventory list failed ({exc})")
        listed = []
    for raw in listed:
        if isinstance(raw, dict):
            merge_switch(store, _as_switch_record(raw))
    print(f"  Inventory list: {len(listed)} record(s), {len(store)} unique MAC(s)")

    print("  GET /api/v1/orgs/{org_id}/inventory/search?type=switch")
    try:
        searched = client.search_paginated(
            f"/orgs/{org_id}/inventory/search",
            params={"type": "switch", "limit": PAGE_LIMIT},
        )
    except MistAPIError as exc:
        print(f"  Inventory search failed ({exc})")
        searched = []
    member_rows = 0
    for raw in searched:
        if not isinstance(raw, dict):
            continue
        merge_switch(store, _as_switch_record(raw))
        members = raw.get("members") or []
        if not isinstance(members, list):
            continue
        for index, member in enumerate(members):
            if isinstance(member, dict) and member.get("mac"):
                merge_switch(store, _member_record(raw, member, index))
                member_rows += 1
    print(f"  Inventory search: {len(searched)} VC/switch row(s), {member_rows} member(s) merged")

    print("  GET /api/v1/sites/{site_id}/devices?type=switch  (every site)")
    site_device_count = 0
    configs_from_list = 0
    for site in sites:
        site_id = site.get("id")
        if not site_id:
            continue
        try:
            devices = client.get_paginated(
                f"/sites/{site_id}/devices",
                params={"type": "switch", "limit": PAGE_LIMIT},
            )
        except MistAPIError as exc:
            print(f"  Skip site devices {site.get('name')}: {exc}")
            continue
        for device in devices:
            if not isinstance(device, dict):
                continue
            record = _as_switch_record(device, site_id=site_id)
            merge_switch(store, record)
            site_device_count += 1
            mac = norm_mac(record.get("mac"))
            if mac and harvest_port_config(device, mac, configs):
                configs_from_list += 1
        time.sleep(RATE_SLEEP)
    print(
        f"  Site device lists: {site_device_count} switch(es); "
        f"port_config embedded on {configs_from_list}"
    )

    switches = list(store.values())
    assigned = sum(1 for s in switches if s.get("site_id"))
    connected = sum(1 for s in switches if s.get("connected") is True)
    members = sum(1 for s in switches if is_secondary_vc_member(s))
    print(
        f"  Unique switches: {len(switches)}    "
        f"Assigned: {assigned}    Connected: {connected}    VC members: {members}"
    )
    return switches, configs


def fetch_port_configs(
    client: MistClient,
    switches: list[dict[str, Any]],
    already: dict[tuple[str, str], dict[str, Any]],
) -> dict[tuple[str, str], dict[str, Any]]:
    print("\n=== Port config  GET /api/v1/sites/{site_id}/devices/{device_id} ===")
    mapping = dict(already)
    have_mac = {mac for mac, _port in mapping}
    fetched = 0
    skipped = 0
    for sw in switches:
        mac = norm_mac(sw.get("mac"))
        site_id = sw.get("site_id")
        device_id = sw.get("id")
        if not mac or mac in have_mac:
            continue
        if is_secondary_vc_member(sw):
            skipped += 1
            continue
        if not site_id or not device_id:
            skipped += 1
            continue
        try:
            cfg = client.get(f"/sites/{site_id}/devices/{device_id}") or {}
        except MistAPIError:
            skipped += 1
            continue
        fetched += 1
        if isinstance(cfg, dict) and harvest_port_config(cfg, mac, mapping):
            have_mac.add(mac)
        time.sleep(RATE_SLEEP)
    print(f"  Loaded config from {fetched} additional switch(es); skipped {skipped}")
    print(f"  Configured port entries: {len(mapping)}")
    return mapping


def extract_stat_ports(st: dict[str, Any]) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    for key in ("ports", "port_stat", "interfaces"):
        value = st.get(key)
        if isinstance(value, list):
            for port in value:
                if isinstance(port, dict):
                    found.append(port)
        elif isinstance(value, dict):
            for pid, port in value.items():
                if isinstance(port, dict):
                    item = dict(port)
                    item.setdefault("port_id", pid)
                    found.append(item)
    return found


def fetch_org_port_stats(client: MistClient, org_id: str) -> list[dict[str, Any]]:
    print("\n=== Port stats  GET /api/v1/orgs/{org_id}/stats/ports/search?device_type=switch ===")
    try:
        ports = client.search_paginated(
            f"/orgs/{org_id}/stats/ports/search",
            params={"device_type": "switch", "limit": PAGE_LIMIT},
        )
        print(f"  Org port search returned {len(ports)} row(s)")
        return ports
    except MistAPIError as exc:
        print(f"  Org port search unavailable ({exc}). Will use per-site port search.")
        return []


def fetch_site_port_stats(
    client: MistClient,
    sites: list[dict[str, Any]],
    only_site_ids: set[str] | None = None,
) -> list[dict[str, Any]]:
    print("\n=== Port stats  GET /api/v1/sites/{site_id}/stats/ports/search?device_type=switch ===")
    rows: list[dict[str, Any]] = []
    for site in sites:
        site_id = site.get("id")
        if not site_id:
            continue
        if only_site_ids is not None and site_id not in only_site_ids:
            continue
        try:
            batch = client.search_paginated(
                f"/sites/{site_id}/stats/ports/search",
                params={"device_type": "switch", "limit": PAGE_LIMIT},
            )
        except MistAPIError as exc:
            print(f"  Skip site ports {site.get('name')}: {exc}")
            continue
        for item in batch:
            if isinstance(item, dict):
                item.setdefault("site_id", site_id)
                rows.append(item)
        time.sleep(RATE_SLEEP)
    print(f"  Site port search returned {len(rows)} row(s)")
    return rows


def fetch_switch_stats(
    client: MistClient,
    org_id: str,
    sites: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    print("\n=== Device stats  GET /api/v1/orgs/{org_id}/stats/devices?type=switch&status=all ===")
    stats: list[dict[str, Any]] = []
    org_ok = False
    try:
        stats = client.get_paginated(
            f"/orgs/{org_id}/stats/devices",
            params={"type": "switch", "status": "all", "limit": PAGE_LIMIT},
        )
        org_ok = True
        print(f"  Org device stats: {len(stats)} switch(es)")
    except MistAPIError as exc:
        print(f"  Org device stats unavailable ({exc}). Falling back to every site.")

    by_mac: dict[str, dict[str, Any]] = {}
    for item in stats:
        if isinstance(item, dict) and norm_mac(item.get("mac")):
            by_mac[norm_mac(item.get("mac"))] = item
    print("  GET /api/v1/sites/{site_id}/stats/devices?type=switch&status=all  (every site)")
    added = 0
    for site in sites:
        site_id = site.get("id")
        if not site_id:
            continue
        try:
            batch = client.get_paginated(
                f"/sites/{site_id}/stats/devices",
                params={"type": "switch", "status": "all", "limit": PAGE_LIMIT},
            )
        except MistAPIError as exc:
            print(f"  Skip site stats {site.get('name')}: {exc}")
            continue
        for item in batch:
            if not isinstance(item, dict):
                continue
            item.setdefault("site_id", site_id)
            mac = norm_mac(item.get("mac"))
            current = by_mac.get(mac) if mac else None
            if current is not None and org_ok:
                site_ports = extract_stat_ports(item)
                if len(site_ports) > len(extract_stat_ports(current)):
                    current["ports"] = site_ports
                if not current.get("site_id"):
                    current["site_id"] = site_id
                continue
            stats.append(item)
            if mac:
                by_mac[mac] = item
            added += 1
        time.sleep(RATE_SLEEP)
    print(f"  Device stats total: {len(stats)} switch(es) ({added} added from sites)")
    return stats


def fetch_missing_device_stats(
    client: MistClient,
    switches: list[dict[str, Any]],
    switch_stats: list[dict[str, Any]],
    ports_by_mac: set[str],
) -> list[dict[str, Any]]:
    known = {norm_mac(s.get("mac")) for s in switch_stats if isinstance(s, dict)}
    need = []
    for sw in switches:
        mac = norm_mac(sw.get("mac"))
        if not mac or mac in ports_by_mac:
            continue
        if is_secondary_vc_member(sw):
            continue
        if not sw.get("site_id") or not sw.get("id"):
            continue
        need.append(sw)
    if not need:
        return switch_stats
    print(f"\n=== Filling {len(need)} switch(es) with no ports via per-device stats ===")
    for sw in need:
        site_id = sw.get("site_id")
        device_id = sw.get("id")
        try:
            item = client.get(f"/sites/{site_id}/stats/devices/{device_id}") or {}
        except MistAPIError as exc:
            print(f"  No stats for {sw.get('name') or sw.get('mac')}: {exc}")
            continue
        if isinstance(item, dict):
            item.setdefault("site_id", site_id)
            item.setdefault("mac", sw.get("mac"))
            mac = norm_mac(item.get("mac"))
            if mac and mac in known:
                current = next(s for s in switch_stats if norm_mac(s.get("mac")) == mac)
                ports = extract_stat_ports(item)
                if len(ports) > len(extract_stat_ports(current)):
                    current["ports"] = ports
            else:
                switch_stats.append(item)
                if mac:
                    known.add(mac)
        time.sleep(RATE_SLEEP)
    return switch_stats


def expand_port_key(key: str) -> list[str]:
    """Expand Mist port_config keys such as 'ge-0/0/0-3' or 'ge-0/0/0,ge-0/0/10'."""
    ports: list[str] = []
    for part in str(key).split(","):
        part = part.strip()
        if not part:
            continue
        full = re.match(r"^([A-Za-z]+-[\d/]+)-([A-Za-z]+-[\d/]+)$", part)
        if full:
            start_name, end_name = full.group(1), full.group(2)
            start_prefix, start_last = start_name.rsplit("/", 1)
            end_prefix, end_last = end_name.rsplit("/", 1)
            if start_prefix == end_prefix and start_last.isdigit() and end_last.isdigit():
                width = len(start_last)
                for i in range(int(start_last), int(end_last) + 1):
                    ports.append(f"{start_prefix}/{i:0{width}d}" if width > 1 else f"{start_prefix}/{i}")
                continue
        if "/" in part:
            prefix, last = part.rsplit("/", 1)
            if "-" in last:
                start_s, end_s = last.split("-", 1)
                if start_s.isdigit() and end_s.isdigit():
                    width = len(start_s)
                    for i in range(int(start_s), int(end_s) + 1):
                        ports.append(f"{prefix}/{i:0{width}d}" if width > 1 else f"{prefix}/{i}")
                    continue
        ports.append(part)
    return ports or [key]


def fpc_from_port(port_id: str) -> int | None:
    match = re.match(r"^[a-z]+-(\d+)/", norm_port(port_id))
    if not match:
        return None
    return int(match.group(1))


def _first(obj: dict[str, Any], *keys: str, default: Any = "") -> Any:
    for key in keys:
        if key in obj and obj[key] not in (None, ""):
            return obj[key]
    return default


def _fmt_speed_duplex(speed: Any, full_duplex: Any, duplex: Any = None) -> str:
    if speed in (None, "", 0, "0"):
        speed_txt = ""
    else:
        try:
            speed_i = int(speed)
            if speed_i >= 1000 and speed_i % 1000 == 0:
                speed_txt = f"{speed_i // 1000}G"
            else:
                speed_txt = f"{speed_i}M"
        except (TypeError, ValueError):
            speed_txt = str(speed)

    if duplex in ("full", "half"):
        dup_txt = duplex
    elif full_duplex is True:
        dup_txt = "full"
    elif full_duplex is False:
        dup_txt = "half"
    else:
        dup_txt = ""

    if speed_txt and dup_txt:
        return f"{speed_txt}/{dup_txt}"
    return speed_txt or dup_txt


def _port_status(up: Any, disabled: Any) -> str:
    if disabled is True:
        return "disabled"
    if up is True:
        return "up"
    if up is False:
        return "down"
    return ""


def _index_switches(
    switches: list[dict[str, Any]],
) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    by_mac: dict[str, dict[str, Any]] = {}
    by_serial: dict[str, dict[str, Any]] = {}
    for sw in switches:
        mac = norm_mac(sw.get("mac"))
        if mac:
            by_mac[mac] = sw
        serial = str(sw.get("serial") or "").upper()
        if serial:
            by_serial[serial] = sw
    for sw in switches:
        vc = norm_mac(sw.get("vc_mac"))
        mac = norm_mac(sw.get("mac"))
        if vc and vc not in by_mac:
            by_mac[vc] = sw
        elif vc and mac == vc:
            by_mac[vc] = sw
    return by_mac, by_serial


def _fpc_members(
    switches: list[dict[str, Any]],
    switch_stats: list[dict[str, Any]],
    by_mac: dict[str, dict[str, Any]],
    by_serial: dict[str, dict[str, Any]],
) -> dict[tuple[str, int], dict[str, Any]]:
    mapping: dict[tuple[str, int], dict[str, Any]] = {}
    for sw in switches:
        mac = norm_mac(sw.get("mac"))
        vc = norm_mac(sw.get("vc_mac"))
        if not mac or not vc or mac == vc:
            continue
        idx = sw.get("member_id")
        if idx is None or idx == "":
            continue
        try:
            fpc = int(idx)
        except (TypeError, ValueError):
            continue
        mapping[(vc, fpc)] = sw

    for st in switch_stats:
        reporting = norm_mac(st.get("mac"))
        vc = norm_mac(st.get("vc_mac")) or reporting
        for mod in st.get("module_stat") or []:
            if not isinstance(mod, dict):
                continue
            idx = mod.get("idx", mod.get("fpc"))
            if idx is None or idx == "":
                continue
            try:
                fpc = int(idx)
            except (TypeError, ValueError):
                continue
            member = None
            serial = str(mod.get("serial") or "").upper()
            if serial:
                member = by_serial.get(serial)
            mmac = norm_mac(mod.get("mac"))
            if member is None and mmac:
                member = by_mac.get(mmac)
            if member is None:
                continue
            if reporting:
                mapping[(reporting, fpc)] = member
            if vc:
                mapping[(vc, fpc)] = member
    return mapping


def build_port_rows(
    sites: list[dict[str, Any]],
    switches: list[dict[str, Any]],
    org_ports: list[dict[str, Any]],
    switch_stats: list[dict[str, Any]],
    port_configs: dict[tuple[str, str], dict[str, Any]],
) -> list[dict[str, Any]]:
    site_by_id = {s.get("id"): s for s in sites if s.get("id")}
    switch_by_mac, by_serial = _index_switches(switches)
    fpc_member = _fpc_members(switches, switch_stats, switch_by_mac, by_serial)

    stats_by_mac: dict[str, dict[str, Any]] = {}
    for st in switch_stats:
        mac = norm_mac(st.get("mac"))
        if mac:
            stats_by_mac[mac] = st

    ports_by_key: dict[tuple[str, str], dict[str, Any]] = {}

    def remember(mac: str, port_id: str, payload: dict[str, Any]) -> None:
        nmac = norm_mac(mac)
        npid = norm_port(port_id)
        if not nmac or not npid:
            return
        key = (nmac, npid)
        existing = ports_by_key.get(key, {})
        merged = {**existing, **{k: v for k, v in payload.items() if v not in (None, "")}}
        merged["mac"] = nmac
        merged["port_id"] = payload.get("port_id") or existing.get("port_id") or npid
        ports_by_key[key] = merged

    for port in org_ports:
        if not isinstance(port, dict):
            continue
        mac = norm_mac(port.get("mac") or port.get("device_mac"))
        port_id = port.get("port_id") or port.get("port") or ""
        if mac and port_id:
            remember(mac, port_id, port)

    for st in switch_stats:
        if not isinstance(st, dict):
            continue
        mac = norm_mac(st.get("mac"))
        for port in extract_stat_ports(st):
            port_id = port.get("port_id") or port.get("port") or port.get("name") or ""
            if mac and port_id:
                payload = dict(port)
                payload.setdefault("mac", mac)
                payload.setdefault("site_id", st.get("site_id"))
                payload.setdefault("device_name", st.get("name"))
                remember(mac, port_id, payload)

    for (mac, port_id) in port_configs:
        if (mac, port_id) not in ports_by_key:
            remember(mac, port_id, {"mac": mac, "port_id": port_id})

    rows: list[dict[str, Any]] = []
    for (mac, port_id), port in sorted(ports_by_key.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        reporting = switch_by_mac.get(mac) or {}
        st = stats_by_mac.get(mac) or {}
        if not st:
            vc_lookup = norm_mac(reporting.get("vc_mac"))
            if vc_lookup:
                st = stats_by_mac.get(vc_lookup) or {}
        site_id = port.get("site_id") or reporting.get("site_id") or st.get("site_id") or ""
        site = site_by_id.get(site_id) or {}

        display = reporting
        vc_mac = norm_mac(port.get("vc_mac") or reporting.get("vc_mac") or st.get("vc_mac"))
        fpc = fpc_from_port(port_id)
        if fpc is not None:
            member = fpc_member.get((mac, fpc)) or (fpc_member.get((vc_mac, fpc)) if vc_mac else None)
            if member is not None:
                display = member
                vc_mac = vc_mac or norm_mac(member.get("vc_mac")) or mac

        cfg = port_configs.get((norm_mac(display.get("mac")), port_id), {})
        if not cfg:
            cfg = port_configs.get((mac, port_id), {})
        if not cfg and vc_mac:
            cfg = port_configs.get((vc_mac, port_id), {})

        speed = _first(port, "speed", default=_first(cfg, "speed"))
        full_duplex = port.get("full_duplex")
        duplex = _first(port, "duplex", default=_first(cfg, "duplex"))
        up = port.get("up")
        disabled = cfg.get("disabled")
        if disabled is None:
            disabled = port.get("disabled")

        description = (
            _first(port, "description", "port_desc", "desc", "if_descr")
            or cfg.get("description")
            or ""
        )
        usage = _first(port, "port_usage", "usage", "profile") or cfg.get("usage") or ""
        sw_status = _first(st, "status")
        if not sw_status:
            if reporting.get("connected") is True or display.get("connected") is True:
                sw_status = "connected"
            elif reporting.get("connected") is False:
                sw_status = "disconnected"

        rows.append(
            {
                "Site": site.get("name") or "",
                "Site ID": site_id,
                "Switch Name": _first(display, "name", "hostname")
                or _first(st, "name")
                or _first(reporting, "name", "hostname")
                or "",
                "Switch Model": _first(display, "model") or _first(st, "model") or _first(reporting, "model") or "",
                "Switch MAC": norm_mac(display.get("mac")) or mac,
                "VC MAC": vc_mac if vc_mac and vc_mac != norm_mac(display.get("mac")) else "",
                "Switch Serial": _first(display, "serial") or _first(st, "serial") or _first(reporting, "serial") or "",
                "Switch Status": sw_status or "",
                "Firmware": _first(st, "version", default=_first(reporting, "version")) or "",
                "Switch IP": _first(st, "ip") or "",
                "Port": port.get("port_id") or port_id,
                "Port Description": description,
                "Port Usage / Profile": usage,
                "Port Status": _port_status(up, disabled),
                "Admin Disabled": "" if disabled in (None, "") else str(bool(disabled)),
                "Link Up": "" if up in (None, "") else str(bool(up)),
                "Speed (Mbps)": speed if speed not in (None, "") else "",
                "Duplex": "full" if full_duplex is True else ("half" if full_duplex is False else duplex or ""),
                "Speed/Duplex": _fmt_speed_duplex(speed, full_duplex, duplex),
                "Uplink": "" if port.get("uplink") in (None, "") else str(bool(port.get("uplink"))),
                "Mode (access/trunk)": cfg.get("mode") or _first(port, "mode") or "",
                "Port Network / Native VLAN": cfg.get("port_network") or _first(port, "port_network", "vlan") or "",
                "Allowed Networks / VLANs": cfg.get("networks") or _first(port, "networks", "vlans") or "",
                "Port Auth": cfg.get("port_auth") or _first(port, "port_auth", "auth_state") or "",
                "Auth State": _first(port, "auth_state") or "",
                "PoE On": "" if port.get("poe_on") in (None, "") else str(bool(port.get("poe_on"))),
                "PoE Mode": _first(port, "poe_mode") or "",
                "PoE Disabled (config)": ""
                if cfg.get("poe_disabled") in (None, "")
                else str(bool(cfg.get("poe_disabled"))),
                "PoE Power Draw (W)": _first(port, "power_draw", "poe_power_draw") or "",
                "Neighbor System Name": _first(port, "neighbor_system_name") or "",
                "Neighbor Port": _first(port, "neighbor_port_desc", "neighbor_port_id") or "",
                "Neighbor MAC": _first(port, "neighbor_mac") or "",
                "STP State": _first(port, "stp_state") or "",
                "STP Role": _first(port, "stp_role") or "",
                "TX bps": _first(port, "tx_bps") or "",
                "RX bps": _first(port, "rx_bps") or "",
                "TX bytes": _first(port, "tx_bytes") or "",
                "RX bytes": _first(port, "rx_bytes") or "",
                "TX errors": _first(port, "tx_errors") or "",
                "RX errors": _first(port, "rx_errors") or "",
                "Port MAC": _first(port, "port_mac") or "",
            }
        )
    return rows


def port_coverage(rows: list[dict[str, Any]]) -> tuple[set[str], set[str]]:
    owners = {norm_mac(r.get("Switch MAC")) for r in rows if norm_mac(r.get("Switch MAC"))}
    vcs = {norm_mac(r.get("VC MAC")) for r in rows if norm_mac(r.get("VC MAC"))}
    return owners, vcs


def switch_has_ports(sw: dict[str, Any], owners: set[str], vcs: set[str]) -> bool:
    mac = norm_mac(sw.get("mac"))
    vc = norm_mac(sw.get("vc_mac"))
    if mac and (mac in owners or mac in vcs):
        return True
    if vc and (vc in owners or vc in vcs):
        return True
    return False


def autosize(ws, max_width: int = 48) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = 0
        for cell in col:
            val = "" if cell.value is None else str(cell.value)
            longest = max(longest, len(val))
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), max_width)


def style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 22


def write_sheet(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_header(ws)
    autosize(ws)
    if rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        table = Table(displayName=ws.title.replace(" ", "")[:30] + "Tbl", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        try:
            ws.add_table(table)
        except ValueError:
            pass


def write_workbook(
    path: str,
    org: dict[str, str],
    region: dict[str, str],
    me: dict[str, Any],
    sites: list[dict[str, Any]],
    switches: list[dict[str, Any]],
    port_rows: list[dict[str, Any]],
) -> None:
    wb = Workbook()
    port_headers = [
        "Site", "Site ID", "Switch Name", "Switch Model", "Switch MAC", "VC MAC",
        "Switch Serial", "Switch Status", "Firmware", "Switch IP", "Port",
        "Port Description", "Port Usage / Profile", "Port Status", "Admin Disabled",
        "Link Up", "Speed (Mbps)", "Duplex", "Speed/Duplex", "Uplink",
        "Mode (access/trunk)", "Port Network / Native VLAN", "Allowed Networks / VLANs",
        "Port Auth", "Auth State", "PoE On", "PoE Mode", "PoE Disabled (config)",
        "PoE Power Draw (W)", "Neighbor System Name", "Neighbor Port", "Neighbor MAC",
        "STP State", "STP Role", "TX bps", "RX bps", "TX bytes", "RX bytes",
        "TX errors", "RX errors", "Port MAC",
    ]
    ws_ports = wb.active
    ws_ports.title = "Switch Ports"
    write_sheet(ws_ports, port_headers, port_rows)

    site_rows = []
    for site in sites:
        site_rows.append(
            {
                "Site": site.get("name") or "",
                "Site ID": site.get("id") or "",
                "Address": site.get("address") or "",
                "Country": site.get("country_code") or "",
                "Timezone": site.get("timezone") or "",
                "Lat": site.get("lat") if site.get("lat") is not None else "",
                "Lng": site.get("lng") if site.get("lng") is not None else "",
            }
        )
    ws_sites = wb.create_sheet("Sites")
    write_sheet(ws_sites, ["Site", "Site ID", "Address", "Country", "Timezone", "Lat", "Lng"], site_rows)

    site_by_id = {s.get("id"): s for s in sites}
    owners, vcs = port_coverage(port_rows)
    sw_rows = []
    for sw in sorted(switches, key=lambda s: (str(s.get("site_id") or ""), str(s.get("name") or ""), str(s.get("mac") or ""))):
        site = site_by_id.get(sw.get("site_id")) or {}
        mac = norm_mac(sw.get("mac"))
        sw_rows.append(
            {
                "Site": site.get("name") or "",
                "Site ID": sw.get("site_id") or "",
                "Switch Name": sw.get("name") or sw.get("hostname") or "",
                "Model": sw.get("model") or "",
                "MAC": mac,
                "Serial": sw.get("serial") or "",
                "Connected": "" if sw.get("connected") is None else str(bool(sw.get("connected"))),
                "Assigned": "yes" if sw.get("site_id") else "no",
                "VC MAC": norm_mac(sw.get("vc_mac")),
                "VC Role": sw.get("vc_role") or "",
                "Adopted": "" if sw.get("adopted") is None else str(bool(sw.get("adopted"))),
                "Ports In Export": "yes" if switch_has_ports(sw, owners, vcs) else "no",
                "Device ID": sw.get("id") or "",
            }
        )
    ws_sw = wb.create_sheet("Switches")
    write_sheet(
        ws_sw,
        [
            "Site", "Site ID", "Switch Name", "Model", "MAC", "Serial", "Connected",
            "Assigned", "VC MAC", "VC Role", "Adopted", "Ports In Export", "Device ID",
        ],
        sw_rows,
    )

    ws_sum = wb.create_sheet("Summary", 0)
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    email = me.get("email") or ""
    name = " ".join(p for p in (me.get("first_name") or "", me.get("last_name") or "") if p)
    macs_with_ports = {norm_mac(r.get("Switch MAC")) for r in port_rows}
    summary_rows = [
        ("Generated", generated),
        ("API Host", region.get("host")),
        ("Region", region.get("name")),
        ("Organization", org.get("name")),
        ("Org ID", org.get("org_id")),
        ("API Identity", name or email or "(org token)"),
        ("API Email", email),
        ("Sites", len(sites)),
        ("Switches", len(switches)),
        ("VC Members", sum(1 for s in switches if is_secondary_vc_member(s))),
        ("Switches With Ports", len(macs_with_ports)),
        ("Switches Without Port Data", sum(1 for s in sw_rows if s.get("Ports In Export") == "no")),
        ("Switch Ports", len(port_rows)),
        ("Ports Up", sum(1 for r in port_rows if r.get("Link Up") == "True")),
        ("Ports Down", sum(1 for r in port_rows if r.get("Link Up") == "False")),
        ("Ports Disabled", sum(1 for r in port_rows if r.get("Admin Disabled") == "True")),
        ("Ports From Config Only", sum(1 for r in port_rows if r.get("Link Up") == "" and r.get("Port Status") in ("", "disabled"))),
    ]
    ws_sum.append(["Field", "Value"])
    for field, value in summary_rows:
        ws_sum.append([field, value])
    style_header(ws_sum)
    autosize(ws_sum, max_width=60)
    ws_sum.sheet_properties.tabColor = "1F4E79"
    wb.save(path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export a Juniper Mist organization-wide switch port inventory to Excel."
    )
    parser.add_argument("--host", help="API host such as api.mist.com (skips region prompt / auto-detect).")
    parser.add_argument("--token", help="API token. Prefer getpass / MIST_API_TOKEN so the token is not stored in history.")
    parser.add_argument("--no-prompt-token", action="store_true", help="Do not prompt; require --token or MIST_API_TOKEN.")
    parser.add_argument("--org-id", help="Organization UUID. Skips the org picker when the token can see multiple orgs.")
    parser.add_argument("--skip-config", action="store_true", help="Do not fetch per-switch port_config (faster, fewer API calls).")
    parser.add_argument("-o", "--output", help="Output .xlsx path. Default: mist_switch_ports_<org>_<timestamp>.xlsx")
    parser.add_argument("--self-test", action="store_true", help="Run offline pagination and merge checks, then exit.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.self_test:
        return _self_test()

    print("Juniper Mist switch port inventory")
    print("=================================")

    token = prompt_token(args.token, args.no_prompt_token)
    client, region = choose_region(MistClient, token, args.host)
    me = self_check(client)

    orgs = orgs_from_self(me)
    if args.org_id:
        match = next((o for o in orgs if o["org_id"] == args.org_id), None)
        org = match or {"org_id": args.org_id, "name": args.org_id, "roles": []}
        if not match:
            try:
                info = client.get(f"/orgs/{args.org_id}") or {}
                org["name"] = info.get("name") or args.org_id
            except MistAPIError:
                pass
        print(f"\nUsing organization: {org['name']}  ({org['org_id']})")
    else:
        org = choose_org(orgs, client)

    org_id = org["org_id"]
    sites = fetch_sites(client, org_id)
    switches, embedded_configs = fetch_switches(client, org_id, sites)
    for sw in switches:
        if sw.get("site_id"):
            ensure_site(client, sites, sw["site_id"])

    org_ports = fetch_org_port_stats(client, org_id)
    switch_stats = fetch_switch_stats(client, org_id, sites)

    port_macs = {norm_mac(p.get("mac") or p.get("device_mac")) for p in org_ports if isinstance(p, dict)}
    for st in switch_stats:
        if extract_stat_ports(st):
            port_macs.add(norm_mac(st.get("mac")))
    missing_sites = {
        sw.get("site_id")
        for sw in switches
        if sw.get("site_id")
        and norm_mac(sw.get("mac")) not in port_macs
        and not is_secondary_vc_member(sw)
    }
    if not org_ports:
        missing_sites = {s.get("id") for s in sites if s.get("id")}
    if missing_sites:
        org_ports.extend(fetch_site_port_stats(client, sites, only_site_ids=missing_sites))

    if args.skip_config:
        port_configs: dict[tuple[str, str], dict[str, Any]] = {}
        print("\nSkipping per-switch port_config fetch (--skip-config).")
    else:
        port_configs = fetch_port_configs(client, switches, embedded_configs)

    seeded: set[str] = set()
    for port in org_ports:
        if isinstance(port, dict):
            seeded.add(norm_mac(port.get("mac") or port.get("device_mac")))
    for st in switch_stats:
        if extract_stat_ports(st):
            seeded.add(norm_mac(st.get("mac")))
    switch_stats = fetch_missing_device_stats(client, switches, switch_stats, seeded)

    print("\n=== Building inventory ===")
    port_rows = build_port_rows(sites, switches, org_ports, switch_stats, port_configs)
    print(f"  Port rows: {len(port_rows)}")
    owners, vcs = port_coverage(port_rows)
    quiet = [sw for sw in switches if not switch_has_ports(sw, owners, vcs)]
    if quiet:
        print(f"  Switches with no port rows: {len(quiet)}")
        for sw in quiet[:30]:
            print(
                f"    - {sw.get('name') or '(unnamed)'}  mac={norm_mac(sw.get('mac'))}  "
                f"site={sw.get('site_id') or 'unassigned'}"
            )
        if len(quiet) > 30:
            print(f"    ... {len(quiet) - 30} more (see Switches sheet, Ports In Export = no)")

    safe_org = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in org.get("name") or "org")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = args.output or f"mist_switch_ports_{safe_org}_{stamp}.xlsx"
    write_workbook(output, org, region, me, sites, switches, port_rows)
    print(f"\nWrote {output}")
    print(f"  Sites={len(sites)}  Switches={len(switches)}  Ports={len(port_rows)}")
    return 0


def _self_test() -> int:
    failures: list[str] = []

    def check(cond: bool, message: str) -> None:
        print(("  ok   " if cond else "  FAIL ") + message)
        if not cond:
            failures.append(message)

    print("Self-test")
    client = MistClient("api.mist.com", "token")
    check(
        client.resolve_url("/api/v1/orgs/abc/stats/ports/search?search_after=CURSOR")
        == "https://api.mist.com/api/v1/orgs/abc/stats/ports/search?search_after=CURSOR",
        "next URL is not prefixed with a second /api/v1",
    )
    check(
        client.resolve_url("/orgs/abc/inventory") == "https://api.mist.com/api/v1/orgs/abc/inventory",
        "API-relative paths still use /api/v1",
    )
    check(
        _encode_params({"vc": True, "unassigned": False}) == {"vc": "true", "unassigned": "false"},
        "booleans encode as true/false",
    )

    pages = {
        1: [{"mac": f"aa{i:04d}", "id": str(i)} for i in range(0, 100)],
        2: [{"mac": f"aa{i:04d}", "id": str(i)} for i in range(100, 200)],
        3: [{"mac": f"aa{i:04d}", "id": str(i)} for i in range(200, 250)],
    }
    calls: list[dict[str, Any]] = []

    class Resp:
        def __init__(self, payload, headers):
            self.status_code = 200
            self._payload = payload
            self.headers = headers
            self.content = b"{}"
            self.text = "{}"

        def json(self):
            return self._payload

    def fake_get(url, params=None, headers=None, timeout=None):
        calls.append({"url": url, "params": dict(params or {})})
        page = int((params or {}).get("page") or 1)
        batch = pages.get(page, [])
        hdrs = {"X-Page-Limit": "1000", "X-Page-Page": str(page), "X-Page-Total": "250"}
        return Resp(batch, hdrs)

    client.session.get = fake_get  # type: ignore[method-assign]
    got = client.get_paginated("/orgs/abc/inventory", {"type": "switch", "vc": True, "limit": 1000})
    check(len(got) == 250, f"list pagination follows X-Page-Total (got {len(got)}, want 250)")
    check(all("page" in c["params"] for c in calls), "list calls send the page query")
    check(calls[0]["params"].get("vc") == "true", "vc=true is lowercase on the wire")

    search_calls: list[str] = []

    def fake_search(url, params=None, headers=None, timeout=None):
        search_calls.append(url if not params else f"{url}|{params}")
        if params and "search_after" in params:
            return Resp({"results": [], "total": 3}, {})
        if "search_after=SECOND" in url:
            return Resp({"results": [{"mac": "m2", "port_id": "ge-0/0/1"}], "total": 3, "next": None}, {})
        if "search_after=FIRST" in url:
            return Resp(
                {
                    "results": [{"mac": "m1", "port_id": "ge-0/0/1"}],
                    "total": 3,
                    "next": "/api/v1/orgs/abc/stats/ports/search?limit=1&device_type=switch&search_after=SECOND",
                },
                {},
            )
        return Resp(
            {
                "results": [{"mac": "m0", "port_id": "ge-0/0/0"}],
                "total": 3,
                "next": "/api/v1/orgs/abc/stats/ports/search?limit=1&device_type=switch&search_after=FIRST",
            },
            {},
        )

    client.session.get = fake_search  # type: ignore[method-assign]
    found = client.search_paginated("/orgs/abc/stats/ports/search", {"device_type": "switch", "limit": 1})
    check(len(found) == 3, f"search pagination follows next URLs (got {len(found)}, want 3)")
    check(
        any(c.startswith("https://api.mist.com/api/v1/orgs/abc/stats/ports/search?limit=1") for c in search_calls),
        "relative next URL is fetched on the API host",
    )

    expanded = expand_port_key("ge-0/0/0-3,xe-0/0/48")
    check(
        expanded == ["ge-0/0/0", "ge-0/0/1", "ge-0/0/2", "ge-0/0/3", "xe-0/0/48"],
        f"port range expansion {expanded}",
    )
    full = expand_port_key("ge-0/0/0-ge-0/0/2")
    check(full == ["ge-0/0/0", "ge-0/0/1", "ge-0/0/2"], f"full-name port range {full}")

    store: dict[str, dict[str, Any]] = {}
    merge_switch(store, _as_switch_record({
        "mac": "020003aabbcc", "name": "closet-vc", "id": "dev-vc", "site_id": "site-1",
        "vc_mac": "020003aabbcc", "model": "EX4400-48P",
    }))
    merge_switch(store, _member_record(
        {"mac": "020003aabbcc", "vc_mac": "020003aabbcc", "name": "closet-vc", "site_id": "site-1"},
        {"mac": "5C:5B:35:00:00:02", "model": "EX4400-48P", "serial": "SN2", "member_id": 1},
        1,
    ))
    check(len(store) == 2, f"VC virtual device and member are both kept ({len(store)})")
    check("5c5b35000002" in store, "member MAC is normalized")
    check(store["5c5b35000002"]["site_id"] == "site-1", "member inherits the VC site")

    sites = [{"id": "site-1", "name": "HQ"}]
    org_ports = [
        {"mac": "020003aabbcc", "port_id": "ge-0/0/1", "up": True, "site_id": "site-1", "speed": 1000},
        {"mac": "020003aabbcc", "port_id": "ge-1/0/4", "up": False, "site_id": "site-1", "speed": 1000},
    ]
    stats = [{
        "mac": "020003aabbcc",
        "name": "closet-vc",
        "site_id": "site-1",
        "module_stat": [
            {"idx": 0, "serial": "SN0", "model": "EX4400-48P"},
            {"idx": 1, "serial": "SN2", "model": "EX4400-48P"},
        ],
        "ports": [
            {"port_id": "ge-0/0/1", "up": True},
            {"port_id": "ge-1/0/4", "up": False},
            {"port_id": "ge-0/0/9", "up": False, "description": "from-stats"},
        ],
    }]
    merge_switch(store, {
        "mac": "aabbccddeef0", "serial": "SN0", "model": "EX4400-48P",
        "vc_mac": "020003aabbcc", "member_id": 0, "site_id": "site-1", "name": "closet-vc fpc0",
    })
    switches = list(store.values())
    configs = {("020003aabbcc", "ge-0/0/7"): {"description": "configured-only", "usage": "access", "disabled": True}}
    rows = build_port_rows(sites, switches, org_ports, stats, configs)
    by_port = {r["Port"]: r for r in rows}
    check(
        set(by_port) >= {"ge-0/0/1", "ge-1/0/4", "ge-0/0/9", "ge-0/0/7"},
        f"stats, search, and config-only ports are all present ({sorted(by_port)})",
    )
    check(by_port["ge-1/0/4"]["Switch MAC"] == "5c5b35000002", "FPC 1 port is attributed to the VC member")
    check(by_port["ge-1/0/4"]["VC MAC"] == "020003aabbcc", "member port keeps the VC MAC")
    check(by_port["ge-0/0/7"]["Port Description"] == "configured-only", "config-only port is exported")
    check(by_port["ge-0/0/7"]["Port Status"] == "disabled", "admin-disabled config port is marked disabled")

    print()
    if failures:
        print(f"{len(failures)} failure(s)")
        return 1
    print("All self-tests passed")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\nCancelled.")
        raise SystemExit(130)
    except MistAPIError as exc:
        print(f"\nError: {exc}", file=sys.stderr)
        raise SystemExit(1)
