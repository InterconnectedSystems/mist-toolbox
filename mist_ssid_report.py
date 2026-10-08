#!/usr/bin/env python3
"""
Juniper Mist (GC2 cloud) - SSID inventory per site -> Excel.

Steps:
  1. Prompt for the API token (getpass, never echoed or stored).
  2. Self check (GET /api/v1/self) to discover which org(s) the token can see.
  3. List every site in the selected org.
  4. For each site, pull the "derived" WLAN list - i.e. every SSID that
     actually applies to that site, whether defined at the site level or
     pushed down from an org-level WLAN template.
  5. Write everything to an .xlsx workbook.

Requirements:  pip install requests openpyxl
"""

import sys
import time
from datetime import datetime
from getpass import getpass

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Mist "Global 04" cloud (GC2).  Change here if you ever need another region.
MIST_API = "https://api.gc2.mist.com/api/v1"

PAGE_LIMIT = 1000
MAX_RETRIES = 5


# --------------------------------------------------------------------------- #
# API helpers
# --------------------------------------------------------------------------- #
def make_session(token: str) -> requests.Session:
    s = requests.Session()
    s.headers.update({
        "Authorization": f"Token {token}",
        "Accept": "application/json",
    })
    return s


def api_get(session: requests.Session, path: str, params=None):
    """GET with basic retry/backoff for 429 rate limits and transient 5xx."""
    url = f"{MIST_API}{path}"
    for attempt in range(1, MAX_RETRIES + 1):
        resp = session.get(url, params=params, timeout=30)
        if resp.status_code == 429 or resp.status_code >= 500:
            wait = int(resp.headers.get("Retry-After", 2 ** attempt))
            print(f"    {resp.status_code} on {path}, retrying in {wait}s "
                  f"({attempt}/{MAX_RETRIES})...")
            time.sleep(wait)
            continue
        if resp.status_code == 401:
            sys.exit("ERROR: 401 Unauthorized - API token is invalid or expired "
                     f"for {MIST_API}.")
        resp.raise_for_status()
        return resp.json()
    resp.raise_for_status()


def api_get_all(session: requests.Session, path: str, params=None):
    """Follow Mist page/limit pagination and return the combined list."""
    params = dict(params or {})
    params["limit"] = PAGE_LIMIT
    page = 1
    results = []
    while True:
        params["page"] = page
        batch = api_get(session, path, params)
        if not batch:
            break
        results.extend(batch)
        if len(batch) < PAGE_LIMIT:
            break
        page += 1
    return results


# --------------------------------------------------------------------------- #
# Discovery
# --------------------------------------------------------------------------- #
def self_check(session: requests.Session):
    """Return (who, orgs) where orgs is a list of {'org_id','name','role'}."""
    me = api_get(session, "/self")
    who = me.get("email") or me.get("name") or "API token"

    orgs = {}
    for priv in me.get("privileges", []):
        org_id = priv.get("org_id")
        if not org_id:
            continue
        # Prefer the org-scope entry for the name/role; site-scope entries
        # still tell us the org exists.
        if priv.get("scope") == "org" or org_id not in orgs:
            orgs[org_id] = {
                "org_id": org_id,
                "name": priv.get("org_name") or priv.get("name") or org_id,
                "role": priv.get("role", ""),
            }
    return who, list(orgs.values())


def choose_org(orgs):
    if not orgs:
        org_id = input("No orgs found in token privileges. Enter Org ID manually: ").strip()
        if not org_id:
            sys.exit("No org selected.")
        return {"org_id": org_id, "name": org_id, "role": ""}

    if len(orgs) == 1:
        return orgs[0]

    print("\nThis token has access to multiple orgs:")
    for i, o in enumerate(orgs, 1):
        print(f"  {i}) {o['name']}  [{o['org_id']}]  role={o['role']}")
    while True:
        choice = input(f"Select org [1-{len(orgs)}]: ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(orgs):
            return orgs[int(choice) - 1]
        print("Invalid choice.")


def get_org_name(session, org):
    """Fetch the real org name in case /self only gave us an ID."""
    try:
        info = api_get(session, f"/orgs/{org['org_id']}")
        return info.get("name") or org["name"]
    except requests.HTTPError:
        return org["name"]


def get_wlan_templates(session, org_id):
    """Map template_id -> template name (used to label where an SSID comes from)."""
    try:
        return {t["id"]: t.get("name", t["id"])
                for t in api_get_all(session, f"/orgs/{org_id}/templates")}
    except requests.HTTPError:
        return {}


# --------------------------------------------------------------------------- #
# WLAN parsing
# --------------------------------------------------------------------------- #
def describe_vlan(wlan):
    if not wlan.get("vlan_enabled"):
        return "untagged"
    if wlan.get("vlan_ids"):
        return ",".join(str(v) for v in wlan["vlan_ids"])
    return str(wlan.get("vlan_id", "")) or "dynamic"


def describe_auth(wlan):
    auth = wlan.get("auth") or {}
    a_type = auth.get("type", "")
    pairwise = auth.get("pairwise") or []
    if pairwise:
        return f"{a_type} ({'/'.join(pairwise)})"
    return a_type


def wlan_source(wlan, templates):
    tid = wlan.get("template_id")
    if tid:
        return f"Org template: {templates.get(tid, tid)}"
    if wlan.get("site_id"):
        return "Site"
    return "Org"


def wlan_row(site, wlan, templates):
    bands = wlan.get("bands") or ([wlan["band"]] if wlan.get("band") else [])
    return [
        site.get("name", ""),
        site.get("id", ""),
        wlan.get("ssid", ""),
        "Yes" if wlan.get("enabled", True) else "No",
        "Yes" if wlan.get("hide_ssid") else "No",
        describe_auth(wlan),
        describe_vlan(wlan),
        ", ".join(str(b) for b in bands),
        wlan.get("interface", ""),
        wlan_source(wlan, templates),
        wlan.get("id", ""),
    ]


# --------------------------------------------------------------------------- #
# Excel output
# --------------------------------------------------------------------------- #
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 60)


def build_workbook(org_name, org_id, detail_rows, summary_rows, errors):
    wb = Workbook()

    ws = wb.active
    ws.title = "SSIDs by Site"
    write_sheet(ws, [
        "Site", "Site ID", "SSID", "Enabled", "Hidden", "Auth",
        "VLAN", "Bands", "Interface", "Source", "WLAN ID",
    ], detail_rows)

    write_sheet(wb.create_sheet("Site Summary"), [
        "Site", "Site ID", "SSID Count", "Enabled SSIDs", "Disabled SSIDs",
    ], summary_rows)

    info = wb.create_sheet("Info")
    for row in [
        ["Org", org_name],
        ["Org ID", org_id],
        ["Cloud", MIST_API],
        ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Sites", len(summary_rows)],
        ["SSID rows", len(detail_rows)],
    ]:
        info.append(row)
    if errors:
        info.append([])
        info.append(["Sites with errors"])
        for site_name, err in errors:
            info.append([site_name, err])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 60

    return wb


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main():
    print(f"Juniper Mist SSID report  ({MIST_API})")
    token = getpass("Mist API token: ").strip()
    if not token:
        sys.exit("No token entered.")

    session = make_session(token)

    print("Running self check...")
    who, orgs = self_check(session)
    print(f"  Authenticated as: {who}")
    org = choose_org(orgs)
    org_id = org["org_id"]
    org_name = get_org_name(session, org)
    print(f"  Org: {org_name}  [{org_id}]")

    print("Fetching sites...")
    sites = sorted(api_get_all(session, f"/orgs/{org_id}/sites"),
                   key=lambda s: s.get("name", "").lower())
    print(f"  Found {len(sites)} site(s)")

    templates = get_wlan_templates(session, org_id)

    detail_rows, summary_rows, errors = [], [], []
    for i, site in enumerate(sites, 1):
        name = site.get("name", site["id"])
        print(f"  [{i}/{len(sites)}] {name}")
        try:
            wlans = api_get(session, f"/sites/{site['id']}/wlans/derived")
        except requests.HTTPError as e:
            print(f"      ! failed: {e}")
            errors.append((name, str(e)))
            continue

        wlans = sorted(wlans, key=lambda w: w.get("ssid", "").lower())
        for w in wlans:
            detail_rows.append(wlan_row(site, w, templates))

        enabled = [w.get("ssid", "") for w in wlans if w.get("enabled", True)]
        disabled = [w.get("ssid", "") for w in wlans if not w.get("enabled", True)]
        summary_rows.append([name, site["id"], len(wlans),
                             ", ".join(enabled), ", ".join(disabled)])

    safe_org = "".join(c if c.isalnum() or c in "-_" else "_" for c in org_name)
    out_file = f"mist_ssids_{safe_org}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    build_workbook(org_name, org_id, detail_rows, summary_rows, errors).save(out_file)

    print(f"\nDone. {len(detail_rows)} SSID rows across {len(summary_rows)} site(s).")
    if errors:
        print(f"{len(errors)} site(s) had errors - see the 'Info' sheet.")
    print(f"Saved: {out_file}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit("\nCancelled.")
    except requests.RequestException as e:
        sys.exit(f"ERROR: {e}")
