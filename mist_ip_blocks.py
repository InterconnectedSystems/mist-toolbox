#!/usr/bin/env python3
"""
mist_irb_report.py

Connects to Juniper Mist, pulls the Mist-generated Junos configuration for
every switch in an org, extracts all IRB interface addresses, computes the
network for each one, and writes everything to an Excel workbook.

Workflow:
  1. Pick the Mist cloud region (or enter a custom API host)
  2. Enter the API token (hidden input via getpass)
  3. Self check: GET /api/v1/self, discover orgs, confirm org access
  4. List all sites in the org
  5. List all switches in the org (inventory, Virtual Chassis aware)
  6. Pull each switch's config (GET .../devices/{id}/config_cmd)
  7. Parse "interfaces irb unit N ..." plus VLAN and routing-instance mappings
  8. Compute network / mask / broadcast / usable range for every address
  9. Write an .xlsx with Summary, IRB Interfaces, Networks, Switches, Sites

Requirements:
  pip install requests openpyxl

Usage:
  python mist_irb_report.py
  python mist_irb_report.py --output irb.xlsx --workers 8 --save-configs configs
  python mist_irb_report.py --host api.eu.mist.com --org-id <org_uuid>
"""

import argparse
import getpass
import ipaddress
import re
import shlex
import sys
import time
from collections import OrderedDict, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

try:
    import requests
    from requests.adapters import HTTPAdapter
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    sys.exit(f"Missing dependency ({exc.name}). Install with: pip install requests openpyxl")


# (label, portal URL you log in to, API host)
MIST_CLOUDS = [
    ("Global 01", "manage.mist.com", "api.mist.com"),
    ("Global 02", "manage.gc1.mist.com", "api.gc1.mist.com"),
    ("Global 03", "manage.ac2.mist.com", "api.ac2.mist.com"),
    ("Global 04", "manage.gc2.mist.com", "api.gc2.mist.com"),
    ("Global 05", "manage.gc4.mist.com", "api.gc4.mist.com"),
    ("EMEA 01", "manage.eu.mist.com", "api.eu.mist.com"),
    ("EMEA 02", "manage.gc3.mist.com", "api.gc3.mist.com"),
    ("EMEA 03", "manage.ac6.mist.com", "api.ac6.mist.com"),
    ("EMEA 04", "manage.gc6.mist.com", "api.gc6.mist.com"),
    ("APAC 01", "manage.ac5.mist.com", "api.ac5.mist.com"),
    ("APAC 02", "manage.gc5.mist.com", "api.gc5.mist.com"),
    ("APAC 03", "manage.gc7.mist.com", "api.gc7.mist.com"),
    ("US Gov", "manage.us.mist-federal.com", "api.us.mist-federal.com"),
]

PAGE_LIMIT = 1000


# --------------------------------------------------------------------------
# Mist API client
# --------------------------------------------------------------------------

class MistAPIError(Exception):
    def __init__(self, status, message, url=""):
        super().__init__(f"HTTP {status}: {message}" if status else message)
        self.status = status
        self.url = url


def _error_text(resp):
    try:
        body = resp.json()
        if isinstance(body, dict):
            return str(body.get("detail") or body.get("message") or body)[:300]
    except ValueError:
        pass
    return (resp.text or resp.reason or "").strip()[:300]


class MistClient:
    def __init__(self, host, token, timeout=60, max_retries=5, pool_size=10):
        self.host = host
        self.base = f"https://{host}"
        self.timeout = timeout
        self.max_retries = max_retries
        self.session = requests.Session()
        adapter = HTTPAdapter(pool_connections=pool_size, pool_maxsize=pool_size)
        self.session.mount("https://", adapter)
        self.session.headers.update({
            "Authorization": f"Token {token}",
            "Accept": "application/json",
            "User-Agent": "mist-irb-report/1.0",
        })

    def _request(self, method, path, params=None):
        url = self.base + path
        resp = None
        for attempt in range(self.max_retries + 1):
            try:
                resp = self.session.request(method, url, params=params, timeout=self.timeout)
            except requests.RequestException as exc:
                if attempt >= self.max_retries:
                    raise MistAPIError(None, f"Connection error: {exc}", url)
                time.sleep(min(2 ** attempt, 30))
                continue

            # Rate limited or server-side hiccup: back off and retry
            if resp.status_code == 429 or resp.status_code >= 500:
                if attempt >= self.max_retries:
                    break
                retry_after = resp.headers.get("Retry-After", "")
                wait = int(retry_after) if retry_after.isdigit() else 2 ** (attempt + 1)
                wait = min(wait, 120)
                if resp.status_code == 429:
                    print(f"    rate limited by Mist, waiting {wait}s ...")
                time.sleep(wait)
                continue

            if resp.status_code >= 400:
                raise MistAPIError(resp.status_code, _error_text(resp), url)
            return resp

        raise MistAPIError(resp.status_code, _error_text(resp), url)

    def get(self, path, params=None):
        return self._request("GET", path, params=params).json()

    def get_all(self, path, params=None):
        """GET a list endpoint, following Mist's limit/page pagination."""
        params = dict(params or {})
        params["limit"] = PAGE_LIMIT
        results = []
        first_ids = set()
        page = 1
        while True:
            params["page"] = page
            resp = self._request("GET", path, params=params)
            data = resp.json()
            if isinstance(data, dict) and "results" in data:
                data = data["results"]
            if not isinstance(data, list):
                return data
            if not data:
                break
            # Guard against endpoints that ignore the page parameter
            marker = repr(data[0])
            if marker in first_ids:
                break
            first_ids.add(marker)

            results.extend(data)
            total = resp.headers.get("X-Page-Total")
            if total and total.isdigit() and len(results) >= int(total):
                break
            if len(data) < PAGE_LIMIT:
                break
            page += 1
        return results


# --------------------------------------------------------------------------
# Interactive helpers
# --------------------------------------------------------------------------

def select_region():
    print("\nSelect your Mist cloud (match the portal URL you log in to):")
    for i, (label, portal, api_host) in enumerate(MIST_CLOUDS, 1):
        print(f"  {i:>2}) {label:<10} {portal:<28} -> {api_host}")
    custom = len(MIST_CLOUDS) + 1
    print(f"  {custom:>2}) Custom API host")

    while True:
        choice = input(f"Choice [1-{custom}, default 1]: ").strip() or "1"
        if choice.isdigit() and 1 <= int(choice) <= custom:
            n = int(choice)
            if n < custom:
                return MIST_CLOUDS[n - 1][2]
            host = input("API host (e.g. api.mist.com): ").strip()
            host = re.sub(r"^https?://", "", host).split("/")[0]
            if host.startswith("manage."):
                host = "api." + host[len("manage."):]
            if host:
                return host
        print("  Invalid choice, try again.")


def discover_orgs(api, me):
    """Build the list of orgs this token can reach from /self privileges."""
    orgs = OrderedDict()
    msp_roles = {}

    for priv in me.get("privileges", []) or []:
        scope = priv.get("scope")
        org_id = priv.get("org_id")
        if scope == "org" and org_id:
            orgs[org_id] = {
                "id": org_id,
                "name": priv.get("org_name") or priv.get("name") or "",
                "role": priv.get("role", ""),
                "scope": "org",
            }
        elif scope == "site" and org_id and org_id not in orgs:
            orgs[org_id] = {
                "id": org_id,
                "name": priv.get("org_name") or "",
                "role": priv.get("role", ""),
                "scope": "site",
            }
        elif scope == "msp" and priv.get("msp_id"):
            msp_roles[priv["msp_id"]] = priv.get("role", "")

    for msp_id, role in msp_roles.items():
        try:
            for org in api.get_all(f"/api/v1/msps/{msp_id}/orgs"):
                if org.get("id") and org["id"] not in orgs:
                    orgs[org["id"]] = {"id": org["id"], "name": org.get("name", ""),
                                       "role": role, "scope": "msp"}
        except MistAPIError as exc:
            print(f"[WARN] Could not list orgs for MSP {msp_id}: {exc}")

    return list(orgs.values())


def choose_org(orgs):
    if not orgs:
        org_id = input("No orgs found in the token's privileges. Enter the org ID: ").strip()
        if not org_id:
            sys.exit("No org selected.")
        return {"id": org_id, "name": "", "role": "", "scope": "manual"}

    if len(orgs) == 1:
        return orgs[0]

    print("\nThis token has access to multiple orgs:")
    for i, org in enumerate(orgs, 1):
        print(f"  {i:>3}) {org['name'] or '(unnamed)':<40} {org['id']}  [{org['scope']}:{org['role']}]")
    while True:
        choice = input(f"Select org [1-{len(orgs)}]: ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(orgs):
            return orgs[int(choice) - 1]
        print("  Invalid choice, try again.")


def self_check(api, org_id_arg):
    print("\n== Self check ==")
    try:
        me = api.get("/api/v1/self")
    except MistAPIError as exc:
        if exc.status == 401:
            sys.exit(f"[FAIL] Token rejected (401). Check the token and that it belongs to "
                     f"the selected cloud ({api.host}).")
        if exc.status is None:
            sys.exit(f"[FAIL] Could not reach {api.host}: {exc}")
        print(f"[WARN] /api/v1/self failed: {exc}")
        me = {}
    else:
        print(f"[ OK ] Connected to {api.host}")

    who = me.get("email") or me.get("name") or "API token"
    if me:
        print(f"[ OK ] Authenticated as: {who}")

    if org_id_arg:
        org = {"id": org_id_arg, "name": "", "role": "", "scope": "cli"}
    else:
        org = choose_org(discover_orgs(api, me))

    try:
        org_info = api.get(f"/api/v1/orgs/{org['id']}")
    except MistAPIError as exc:
        sys.exit(f"[FAIL] Cannot read org {org['id']}: {exc}")
    org["name"] = org_info.get("name") or org["name"]

    role = f", role: {org['role']}" if org.get("role") else ""
    print(f"[ OK ] Org access: {org['name']} ({org['id']}){role}")
    if org.get("scope") == "site":
        print("[WARN] Token only has site-level privileges in this org; org-wide "
              "inventory calls may fail.")
    return who, org


# --------------------------------------------------------------------------
# Inventory
# --------------------------------------------------------------------------

def build_switch_list(inventory, sites_by_id):
    """Collapse inventory entries into one record per managed switch / VC."""
    groups = OrderedDict()
    for item in inventory:
        mac = (item.get("mac") or "").lower()
        key = (item.get("vc_mac") or mac).lower()
        if key:
            groups.setdefault(key, []).append(item)

    switches = []
    for key, members in groups.items():
        primary = next((m for m in members if (m.get("mac") or "").lower() == key), members[0])
        device_id = primary.get("id") if (primary.get("mac") or "").lower() == key else None
        device_id = device_id or f"00000000-0000-0000-1000-{key}"
        site_id = next((m.get("site_id") for m in members if m.get("site_id")), None)
        name = (primary.get("name") or primary.get("hostname")
                or next((m.get("name") for m in members if m.get("name")), "") or key)
        models = sorted({m.get("model") for m in members if m.get("model")})
        serials = [m.get("serial") for m in members if m.get("serial")]

        switches.append({
            "name": name,
            "mac": key,
            "model": ", ".join(models),
            "serials": ", ".join(serials),
            "members": len(members),
            "connected": any(m.get("connected") for m in members),
            "device_id": device_id,
            "site_id": site_id,
            "site_name": (sites_by_id.get(site_id) or {}).get("name", "") if site_id else "",
            "status": "",
            "error": "",
            "cli_lines": 0,
            "irb_units": 0,
            "addresses": 0,
        })

    switches.sort(key=lambda s: (s["site_name"].lower(), s["name"].lower()))
    return switches


# --------------------------------------------------------------------------
# Config parsing
# --------------------------------------------------------------------------

IGNORED_VERBS = {"delete", "deactivate", "activate", "insert", "rename",
                 "annotate", "protect", "unprotect"}
DHCP_KEYWORDS = {"dhcp", "dhcp-client", "dhcpv6-client", "autoconfig"}
VIRTUAL_ADDR_KEYWORDS = {"virtual-gateway-address", "virtual-address",
                         "virtual-inet6-address", "virtual-link-local-address"}
FAMILY_LABEL = {"inet": "IPv4", "inet6": "IPv6"}


def _tokenize(line):
    try:
        return shlex.split(line)
    except ValueError:
        return line.split()


def _handle_irb_unit(unit, rest):
    if not rest:
        return
    key = rest[0]
    if key == "description":
        unit["description"] = " ".join(rest[1:])
    elif key == "disable":
        unit["disabled"] = True
    elif key == "family" and len(rest) >= 3:
        family, frest = rest[1], rest[2:]
        if frest[0] == "address" and len(frest) >= 2:
            addr = unit["addresses"].setdefault(
                (family, frest[1]),
                {"family": family, "address": frest[1], "flags": [], "virtual": []},
            )
            attrs = frest[2:]
            for i, tok in enumerate(attrs):
                nxt = attrs[i + 1] if i + 1 < len(attrs) else None
                if tok in ("primary", "preferred") and tok not in addr["flags"]:
                    addr["flags"].append(tok)
                elif tok in VIRTUAL_ADDR_KEYWORDS and nxt and nxt not in addr["virtual"]:
                    addr["virtual"].append(nxt)
                elif tok in ("vrrp-group", "vrrp-inet6-group") and nxt:
                    flag = f"vrrp-group {nxt}"
                    if flag not in addr["flags"]:
                        addr["flags"].append(flag)
        elif frest[0] in DHCP_KEYWORDS:
            entry = (family, frest[0])
            if entry not in unit["dhcp"]:
                unit["dhcp"].append(entry)


def parse_irb_config(cli_lines):
    """
    Parse Junos 'set' commands and return IRB units plus the VLAN and
    routing-instance mappings that reference them.  Handles statements
    wrapped in 'groups <name>' (Mist puts its config in 'groups top').
    """
    units = OrderedDict()
    vlans = defaultdict(dict)
    ri_by_iface = {}

    def handle_vlan(name, rest):
        if len(rest) < 2:
            return
        if rest[0] == "vlan-id":
            vlans[name]["vlan_id"] = rest[1]
        elif rest[0] == "l3-interface":
            vlans[name]["l3_interface"] = rest[1]

    for raw in cli_lines:
        line = str(raw).strip()
        if not line or line.startswith("#"):
            continue
        tok = _tokenize(line)
        if not tok or tok[0] in IGNORED_VERBS:
            continue
        if tok[0] == "set":
            tok = tok[1:]
        if len(tok) >= 2 and tok[0] == "groups":
            tok = tok[2:]
        if len(tok) < 3:
            continue

        if tok[0] == "interfaces":
            if tok[1] == "irb" and tok[2] == "unit" and len(tok) >= 4:
                unit_no, rest = tok[3], tok[4:]
            elif tok[1].startswith("irb."):
                unit_no, rest = tok[1].split(".", 1)[1], tok[2:]
            else:
                continue
            unit = units.setdefault(unit_no, {"description": "", "disabled": False,
                                              "addresses": OrderedDict(), "dhcp": []})
            _handle_irb_unit(unit, rest)

        elif tok[0] == "vlans":
            handle_vlan(tok[1], tok[2:])

        elif tok[0] == "routing-instances":
            ri, rest = tok[1], tok[2:]
            if rest[0] == "interface" and len(rest) > 1 and rest[1].startswith("irb."):
                ri_by_iface[rest[1]] = ri
            elif rest[0] == "vlans" and len(rest) > 2:
                handle_vlan(rest[1], rest[2:])

    return {"units": units, "vlans": vlans, "routing_instances": ri_by_iface}


# --------------------------------------------------------------------------
# Network math
# --------------------------------------------------------------------------

def network_details(cidr):
    """Compute network facts for an interface address like 10.1.1.1/24."""
    try:
        iface = ipaddress.ip_interface(cidr)
    except ValueError:
        return {"Notes": f"Unparseable address '{cidr}'"}

    ip, net = iface.ip, iface.network
    notes = []
    d = {
        "IP Address": str(ip),
        "Prefix Length": net.prefixlen,
        "Network": str(net),
        "Network Address": str(net.network_address),
        "_network": net,
    }

    if net.version == 4:
        d["Subnet Mask"] = str(net.netmask)
        d["Wildcard Mask"] = str(net.hostmask)
        if net.prefixlen <= 30:
            d["Broadcast"] = str(net.broadcast_address)
            d["First Usable"] = str(net.network_address + 1)
            d["Last Usable"] = str(net.broadcast_address - 1)
            d["Usable Hosts"] = net.num_addresses - 2
            if ip == net.network_address:
                notes.append("IP is the network address")
            elif ip == net.broadcast_address:
                notes.append("IP is the broadcast address")
        elif net.prefixlen == 31:
            d["Broadcast"] = "n/a (/31)"
            d["First Usable"] = str(net[0])
            d["Last Usable"] = str(net[1])
            d["Usable Hosts"] = 2
        else:
            d["Broadcast"] = "n/a (/32)"
            d["First Usable"] = d["Last Usable"] = str(ip)
            d["Usable Hosts"] = 1
    else:
        host_bits = 128 - net.prefixlen
        d["Broadcast"] = "n/a (IPv6)"
        d["First Usable"] = str(net.network_address)
        d["Last Usable"] = str(net.broadcast_address)
        d["Usable Hosts"] = net.num_addresses if host_bits <= 32 else f"2^{host_bits}"
        if ip.is_link_local:
            notes.append("Link-local")

    d["_notes"] = notes
    return d


def _unit_sort_key(unit_no):
    return (0, int(unit_no), "") if unit_no.isdigit() else (1, 0, unit_no)


def _maybe_int(value):
    return int(value) if isinstance(value, str) and value.isdigit() else value


def irb_rows_for_switch(sw, parsed):
    l3_to_vlan = {}
    for name, v in parsed["vlans"].items():
        if v.get("l3_interface"):
            l3_to_vlan.setdefault(v["l3_interface"], (name, v.get("vlan_id", "")))

    rows = []
    for unit_no in sorted(parsed["units"], key=_unit_sort_key):
        unit = parsed["units"][unit_no]
        ifname = f"irb.{unit_no}"
        vlan_name, vlan_id = l3_to_vlan.get(ifname, ("", ""))
        base = {
            "Site": sw["site_name"],
            "Switch": sw["name"],
            "Switch MAC": sw["mac"],
            "Model": sw["model"],
            "Interface": ifname,
            "VLAN Name": vlan_name,
            "VLAN ID": _maybe_int(vlan_id),
            "Routing Instance": parsed["routing_instances"].get(ifname, ""),
            "Description": unit["description"],
            "Site ID": sw["site_id"],
            "Device ID": sw["device_id"],
        }
        unit_flags = ["disabled"] if unit["disabled"] else []

        for addr in unit["addresses"].values():
            row = dict(base)
            row["Family"] = FAMILY_LABEL.get(addr["family"], addr["family"])
            row["Address Type"] = "static"
            row["Interface Address"] = addr["address"]
            details = network_details(addr["address"])
            notes = details.pop("_notes", [])
            if "Notes" in details:
                notes.append(details.pop("Notes"))
            row.update(details)
            row["Virtual / VRRP Address"] = ", ".join(addr["virtual"])
            row["Flags"] = ", ".join(unit_flags + addr["flags"])
            row["Notes"] = "; ".join(notes)
            rows.append(row)

        for family, keyword in unit["dhcp"]:
            row = dict(base)
            row["Family"] = FAMILY_LABEL.get(family, family)
            row["Address Type"] = keyword
            row["Flags"] = ", ".join(unit_flags)
            row["Notes"] = "Address assigned dynamically; not in config"
            rows.append(row)

        if not unit["addresses"] and not unit["dhcp"]:
            row = dict(base)
            row["Address Type"] = "none"
            row["Flags"] = ", ".join(unit_flags)
            row["Notes"] = "IRB unit has no address configured"
            rows.append(row)

    return rows


def summarize_networks(iface_rows):
    nets = OrderedDict()
    for row in iface_rows:
        net = row.get("_network")
        if net is None:
            continue
        agg = nets.setdefault(net, {
            "addresses": [], "switches": set(), "sites": set(),
            "vlan_ids": set(), "vlan_names": set(), "ris": set(),
        })
        agg["addresses"].append(row["IP Address"])
        agg["switches"].add(f"{row['Site']}/{row['Switch']}")
        agg["sites"].add(row["Site"] or row["Site ID"] or "")
        if row.get("VLAN ID") not in ("", None):
            agg["vlan_ids"].add(str(row["VLAN ID"]))
        if row.get("VLAN Name"):
            agg["vlan_names"].add(row["VLAN Name"])
        if row.get("Routing Instance"):
            agg["ris"].add(row["Routing Instance"])

    # Overlap detection: CIDR blocks are either disjoint or nested, so a sweep
    # over networks sorted by start address only needs the still-open blocks.
    overlaps = defaultdict(set)
    for version in (4, 6):
        ordered = sorted((n for n in nets if n.version == version and not n.is_link_local),
                         key=lambda n: (int(n.network_address), n.prefixlen))
        open_blocks = []
        for net in ordered:
            start = int(net.network_address)
            open_blocks = [o for o in open_blocks if int(o.broadcast_address) >= start]
            for other in open_blocks:
                overlaps[net].add(other)
                overlaps[other].add(net)
            open_blocks.append(net)

    rows = []
    for net in sorted(nets, key=lambda n: (n.version, int(n.network_address), n.prefixlen)):
        agg = nets[net]
        details = network_details(str(net))
        ov = sorted(overlaps.get(net, ()), key=lambda n: (int(n.network_address), n.prefixlen))
        rows.append({
            "Network": str(net),
            "Family": f"IPv{net.version}",
            "Prefix Length": net.prefixlen,
            "Subnet Mask": details.get("Subnet Mask", ""),
            "Usable Hosts": details.get("Usable Hosts", ""),
            "Interface Count": len(agg["addresses"]),
            "Switch Count": len(agg["switches"]),
            "Site Count": len(agg["sites"]),
            "Seen At Multiple Sites": "Yes" if len(agg["sites"]) > 1 else "No",
            "Overlaps With": ", ".join(str(n) for n in ov[:25]) + (" ..." if len(ov) > 25 else ""),
            "VLAN IDs": ", ".join(sorted(agg["vlan_ids"], key=lambda v: (len(v), v))),
            "VLAN Names": ", ".join(sorted(agg["vlan_names"])),
            "Routing Instances": ", ".join(sorted(agg["ris"])),
            "Sites": ", ".join(sorted(agg["sites"])),
            "Switches": ", ".join(sorted(agg["switches"])),
            "Interface IPs": ", ".join(agg["addresses"]),
        })
    return rows, sum(1 for n in nets if overlaps.get(n))


# --------------------------------------------------------------------------
# Collection
# --------------------------------------------------------------------------

def _safe_filename(text):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_") or "unnamed"


def collect(api, org_id, workers=8, save_dir=None):
    print("\n== Sites ==")
    sites = api.get_all(f"/api/v1/orgs/{org_id}/sites")
    sites_by_id = {s["id"]: s for s in sites if s.get("id")}
    print(f"Found {len(sites)} site(s)")

    print("\n== Switches ==")
    inventory = api.get_all(f"/api/v1/orgs/{org_id}/inventory", {"type": "switch", "vc": "true"})
    switches = build_switch_list(inventory, sites_by_id)
    assigned = [s for s in switches if s["site_id"]]
    print(f"Found {len(inventory)} switch inventory record(s) -> {len(switches)} switch(es) / "
          f"virtual chassis, {len(assigned)} assigned to a site")
    for sw in switches:
        if not sw["site_id"]:
            sw["status"] = "Skipped - not assigned to a site"

    if save_dir:
        save_dir = Path(save_dir)
        save_dir.mkdir(parents=True, exist_ok=True)

    print(f"\n== Pulling configs ({workers} parallel) ==")

    def fetch(sw):
        data = api.get(f"/api/v1/sites/{sw['site_id']}/devices/{sw['device_id']}/config_cmd")
        cli = data.get("cli", []) if isinstance(data, dict) else data
        if isinstance(cli, str):
            cli = cli.splitlines()
        return cli or []

    iface_rows = []
    width = len(str(len(assigned)))
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {pool.submit(fetch, sw): sw for sw in assigned}
        for n, fut in enumerate(as_completed(futures), 1):
            sw = futures[fut]
            label = f"[{n:>{width}}/{len(assigned)}] {sw['site_name'] or sw['site_id']} / {sw['name']}"
            try:
                cli = fut.result()
            except Exception as exc:  # keep going if one switch fails
                sw["status"] = "Error"
                sw["error"] = str(exc)
                print(f"{label}: ERROR {exc}")
                continue

            sw["cli_lines"] = len(cli)
            if not cli:
                sw["status"] = "No config returned"
                print(f"{label}: no config returned")
                continue

            if save_dir:
                fname = _safe_filename(f"{sw['site_name']}__{sw['name']}__{sw['mac']}") + ".txt"
                (save_dir / fname).write_text("\n".join(map(str, cli)) + "\n", encoding="utf-8")

            parsed = parse_irb_config(cli)
            rows = irb_rows_for_switch(sw, parsed)
            iface_rows.extend(rows)
            sw["irb_units"] = len(parsed["units"])
            sw["addresses"] = sum(1 for r in rows if r.get("Address Type") == "static")
            sw["status"] = "OK" if sw["irb_units"] else "OK - no IRB interfaces"
            print(f"{label}: {sw['irb_units']} IRB unit(s), {sw['addresses']} address(es)")

    iface_rows.sort(key=lambda r: (str(r["Site"]).lower(), str(r["Switch"]).lower(),
                                   _unit_sort_key(r["Interface"].split(".", 1)[1]),
                                   r.get("Family", "")))
    return sites, switches, iface_rows


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")

INTERFACE_COLUMNS = [
    "Site", "Switch", "Switch MAC", "Model", "Interface", "VLAN Name", "VLAN ID",
    "Routing Instance", "Description", "Family", "Address Type", "Interface Address",
    "IP Address", "Prefix Length", "Subnet Mask", "Wildcard Mask", "Network",
    "Network Address", "Broadcast", "First Usable", "Last Usable", "Usable Hosts",
    "Virtual / VRRP Address", "Flags", "Notes", "Site ID", "Device ID",
]
NETWORK_COLUMNS = [
    "Network", "Family", "Prefix Length", "Subnet Mask", "Usable Hosts", "Interface Count",
    "Switch Count", "Site Count", "Seen At Multiple Sites", "Overlaps With", "VLAN IDs",
    "VLAN Names", "Routing Instances", "Sites", "Switches", "Interface IPs",
]
SWITCH_COLUMNS = [
    "Site", "Switch", "MAC", "Model", "Serial(s)", "VC Members", "Connected",
    "Config Status", "Config Lines", "IRB Units", "IRB Addresses", "Error", "Site ID", "Device ID",
]
SITE_COLUMNS = [
    "Site", "Address", "Country", "Timezone", "Switches", "IRB Addresses", "Site ID",
]


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (list, tuple, set)):
        value = ", ".join(map(str, value))
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def add_sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title=title)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")

    widths = [len(c) for c in columns]
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, col in enumerate(columns, start=1):
            value = _cell_value(row.get(col))
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # never let config text become a formula
            if value is not None:
                widths[c_idx - 1] = max(widths[c_idx - 1], len(str(value)))

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 60)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    return ws


def write_workbook(path, meta, sites, switches, iface_rows):
    net_rows, overlap_count = summarize_networks(iface_rows)

    switch_rows = [{
        "Site": s["site_name"], "Switch": s["name"], "MAC": s["mac"], "Model": s["model"],
        "Serial(s)": s["serials"], "VC Members": s["members"], "Connected": s["connected"],
        "Config Status": s["status"], "Config Lines": s["cli_lines"], "IRB Units": s["irb_units"],
        "IRB Addresses": s["addresses"], "Error": s["error"], "Site ID": s["site_id"] or "",
        "Device ID": s["device_id"],
    } for s in switches]

    per_site_switches = defaultdict(int)
    per_site_addrs = defaultdict(int)
    for s in switches:
        if s["site_id"]:
            per_site_switches[s["site_id"]] += 1
            per_site_addrs[s["site_id"]] += s["addresses"]
    site_rows = sorted(({
        "Site": s.get("name", ""), "Address": s.get("address", ""),
        "Country": s.get("country_code", ""), "Timezone": s.get("timezone", ""),
        "Switches": per_site_switches.get(s.get("id"), 0),
        "IRB Addresses": per_site_addrs.get(s.get("id"), 0), "Site ID": s.get("id", ""),
    } for s in sites), key=lambda r: str(r["Site"]).lower())

    static_rows = [r for r in iface_rows if r.get("Address Type") == "static"]
    summary = [
        ("Org", meta["org_name"]),
        ("Org ID", meta["org_id"]),
        ("Mist API host", meta["host"]),
        ("Run by", meta["who"]),
        ("Generated", meta["generated"]),
        ("Sites", len(sites)),
        ("Switches / VCs", len(switches)),
        ("Switches with config pulled", sum(1 for s in switches if s["status"].startswith("OK"))),
        ("Switches with errors", sum(1 for s in switches if s["status"] == "Error")),
        ("Switches not assigned to a site", sum(1 for s in switches if not s["site_id"])),
        ("IRB units", sum(s["irb_units"] for s in switches)),
        ("IRB static addresses", len(static_rows)),
        ("IPv4 addresses", sum(1 for r in static_rows if r.get("Family") == "IPv4")),
        ("IPv6 addresses", sum(1 for r in static_rows if r.get("Family") == "IPv6")),
        ("Unique networks", len(net_rows)),
        ("Networks seen at multiple sites", sum(1 for r in net_rows if r["Seen At Multiple Sites"] == "Yes")),
        ("Networks overlapping another network", overlap_count),
    ]

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Summary", ["Item", "Value"], [{"Item": k, "Value": v} for k, v in summary])
    add_sheet(wb, "IRB Interfaces", INTERFACE_COLUMNS, iface_rows)
    add_sheet(wb, "Networks", NETWORK_COLUMNS, net_rows)
    add_sheet(wb, "Switches", SWITCH_COLUMNS, switch_rows)
    add_sheet(wb, "Sites", SITE_COLUMNS, site_rows)
    wb.save(path)
    return summary


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(description="Export IRB interface addresses from Juniper Mist switches to Excel.")
    p.add_argument("-o", "--output", help="Output .xlsx path (default: mist_irb_<org>_<timestamp>.xlsx)")
    p.add_argument("--host", help="Mist API host, skips the region prompt (e.g. api.eu.mist.com)")
    p.add_argument("--org-id", help="Org ID, skips org discovery/selection")
    p.add_argument("--workers", type=int, default=8, help="Parallel config downloads (default 8)")
    p.add_argument("--save-configs", metavar="DIR", help="Also save each switch's set commands to DIR")
    return p.parse_args()


def main():
    args = parse_args()
    print("Juniper Mist - IRB interface address report")

    host = args.host or select_region()
    token = getpass.getpass(f"Mist API token for {host} (input hidden): ").strip()
    if not token:
        sys.exit("No API token entered.")

    api = MistClient(host, token, pool_size=max(args.workers, 1) + 2)
    who, org = self_check(api, args.org_id)

    try:
        sites, switches, iface_rows = collect(api, org["id"], args.workers, args.save_configs)
    except MistAPIError as exc:
        sys.exit(f"[FAIL] {exc}")

    stamp = datetime.now()
    output = args.output or f"mist_irb_{_safe_filename(org['name'] or org['id'])}_{stamp:%Y%m%d_%H%M%S}.xlsx"
    meta = {"org_name": org["name"], "org_id": org["id"], "host": host, "who": who,
            "generated": stamp.strftime("%Y-%m-%d %H:%M:%S")}

    summary = write_workbook(output, meta, sites, switches, iface_rows)

    print("\n== Summary ==")
    for key, value in summary[5:]:
        print(f"  {key:<38} {value}")
    print(f"\nWrote {Path(output).resolve()}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit("\nCancelled.")
